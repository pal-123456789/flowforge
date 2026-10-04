#!/usr/bin/env python3
"""
FlowForge - ALGOTHON'26 pitch deck (professional edition).
Run:  pip install python-pptx   &&   python make_ppt.py
Output: FlowForge_Deck.pptx   (16:9)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- professional palette ----------
PAPER   = RGBColor(0xFB, 0xFC, 0xFE)   # off-white page
CARD    = RGBColor(0xFF, 0xFF, 0xFF)   # white card
INK     = RGBColor(0x0E, 0x1B, 0x2E)   # deep navy ink (headings)
BODY    = RGBColor(0x45, 0x50, 0x63)   # slate body text
MUTE    = RGBColor(0x8A, 0x93, 0xA6)   # muted caption
HAIR    = RGBColor(0xE4, 0xE8, 0xF0)   # hairline
NAVY    = RGBColor(0x11, 0x24, 0x40)   # dark section bg
ACCENT  = RGBColor(0x3B, 0x5B, 0xDB)   # royal blue accent
ACCENT2 = RGBColor(0x1E, 0xA6, 0x8C)   # teal secondary
GOLD    = RGBColor(0xC2, 0x8A, 0x2B)   # restrained gold
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SOFTB   = RGBColor(0xEE, 0xF2, 0xFC)   # soft blue tint

# category colors (muted, professional)
C_TRIG = RGBColor(0x2E, 0x8B, 0x57)
C_ACT  = RGBColor(0x2B, 0x5C, 0xB8)
C_LOG  = RGBColor(0xB8, 0x7A, 0x1E)
C_DAT  = RGBColor(0xA8, 0x3A, 0x6E)
C_AI   = RGBColor(0x1E, 0x8E, 0xA6)

HEAD = "Georgia"            # serif display for headings
BODYF = "Segoe UI"          # clean sans for body
MONO = "Consolas"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def bg(slide, color=PAPER):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    slide.shapes._spTree.remove(s._element)
    slide.shapes._spTree.insert(2, s._element)
    return s


def rect(slide, x, y, w, h, fill=CARD, line=None, lw=1.0, rounded=False, shadow=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if shadow:
        el = shp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {})
        sh = ef.makeelement(qn('a:outerShdw'),
                            {'blurRad': '90000', 'dist': '30000',
                             'dir': '5400000', 'rotWithShape': '0'})
        clr = sh.makeelement(qn('a:srgbClr'), {'val': '0E1B2E'})
        alpha = clr.makeelement(qn('a:alpha'), {'val': '12000'})
        clr.append(alpha); sh.append(clr); ef.append(sh); el.append(ef)
    return shp


def txt(slide, x, y, w, h, text, size=18, color=INK, bold=False,
        align=PP_ALIGN.LEFT, font=BODYF, anchor=MSO_ANCHOR.TOP, spacing=1.0,
        italic=False, tracking=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = font
        if tracking is not None:
            r.font._rPr.set('spc', str(tracking))
    return tb


def kicker(slide, x, y, text, color=ACCENT):
    txt(slide, x, y, Inches(8), Inches(0.35), text.upper(), size=12.5,
        color=color, bold=True, font=BODYF, tracking=260)
    rect(slide, x+Inches(0.02), y+Inches(0.42), Inches(0.6), Inches(0.03), fill=color)


def page_num(slide, n):
    txt(slide, Inches(12.4), Inches(7.0), Inches(0.8), Inches(0.35),
        f"{n:02d}", size=11, color=MUTE, align=PP_ALIGN.RIGHT, font=MONO)
    txt(slide, Inches(0.9), Inches(7.0), Inches(4), Inches(0.35),
        "FlowForge", size=11, color=MUTE, font=BODYF, tracking=120)


# ===================================================================
# SLIDE 1 - TITLE (light, premium cover)
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 0, Inches(0.18), EMU_H, fill=ACCENT)
rect(s, Inches(0.18), 0, Inches(0.06), EMU_H, fill=ACCENT2)
txt(s, Inches(1.1), Inches(1.05), Inches(10), Inches(0.4),
    "ALGOTHON'26   ·   ALG-AUTO-01   ·   VISUAL WORKFLOW AUTOMATION",
    size=12.5, color=ACCENT, bold=True, tracking=200)
rect(s, Inches(1.12), Inches(1.55), Inches(0.9), Inches(0.035), fill=ACCENT2)
txt(s, Inches(1.05), Inches(2.15), Inches(11.2), Inches(1.6), "FlowForge",
    size=88, color=INK, bold=True, font=HEAD)
txt(s, Inches(1.12), Inches(3.75), Inches(10.8), Inches(1.0),
    "Build, run, and watch real automations \u2014 on a visual canvas, with no code.",
    size=23, color=BODY, font=BODYF, spacing=1.15)
# thin divider
rect(s, Inches(1.12), Inches(5.15), Inches(11.1), Inches(0.015), fill=HAIR)
txt(s, Inches(1.12), Inches(5.4), Inches(11), Inches(0.5),
    "A production-grade workflow platform  ·  20 executable nodes  ·  a from-scratch DAG engine  ·  100% offline",
    size=14.5, color=MUTE, font=BODYF)
txt(s, Inches(1.12), Inches(6.25), Inches(11), Inches(0.4),
    "flowforge-lac-eight.vercel.app", size=13.5, color=ACCENT2, bold=True, font=MONO)

# ===================================================================
# SLIDE 2 - PROBLEM
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, Inches(0.9), Inches(0.7), "The Problem")
txt(s, Inches(0.85), Inches(1.25), Inches(11.6), Inches(1.0),
    "Automation is still locked behind code.", size=38, color=INK, bold=True, font=HEAD)
items = [
    ("Glue code everywhere", "Teams wire APIs together with brittle scripts that only engineers can read or maintain."),
    ("No visibility", "When a job fails overnight, nobody can see which step broke \u2014 or why it broke."),
    ("Expensive & locked-in", "Existing no-code platforms are cloud-locked, paywalled, and require live API keys to even demo."),
]
y = Inches(2.75)
for i, (h, b) in enumerate(items):
    x = Inches(0.9) + i * Inches(4.05)
    card = rect(s, x, y, Inches(3.75), Inches(3.2), fill=CARD, line=HAIR, lw=1.0, rounded=True, shadow=True)
    txt(s, x+Inches(0.35), y+Inches(0.35), Inches(1), Inches(0.6), f"0{i+1}",
        size=30, color=ACCENT, bold=True, font=HEAD)
    rect(s, x+Inches(0.35), y+Inches(1.1), Inches(0.5), Inches(0.03), fill=ACCENT2)
    txt(s, x+Inches(0.35), y+Inches(1.3), Inches(3.1), Inches(0.7), h, size=19, color=INK, bold=True, font=BODYF)
    txt(s, x+Inches(0.35), y+Inches(2.0), Inches(3.1), Inches(1.1), b, size=14, color=BODY, spacing=1.2)
page_num(s, 2)

# ===================================================================
# SLIDE 3 - SOLUTION
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, Inches(0.9), Inches(0.7), "Our Solution", color=ACCENT2)
txt(s, Inches(0.85), Inches(1.25), Inches(11.6), Inches(1.0),
    "Drag. Connect. Run. Watch it happen.", size=38, color=INK, bold=True, font=HEAD)
feats = [
    ("Visual canvas", "An infinite canvas \u2014 drag nodes, connect handles, and branch logic without writing a line of code.", C_ACT),
    ("Real execution", "A topological DAG engine runs the graph live; every node lights up as it executes in real time.", ACCENT2),
    ("Full observability", "Timeline, streaming logs, and per-node input / output for every run, plus an analytics dashboard.", C_DAT),
    ("100% offline", "No database, no paid keys. AI nodes fall back to a deterministic on-device mock so demos never break.", C_TRIG),
]
coords = [(Inches(0.9), Inches(2.7)), (Inches(6.95), Inches(2.7)),
          (Inches(0.9), Inches(4.95)), (Inches(6.95), Inches(4.95))]
for (x, y), (h, b, c) in zip(coords, feats):
    rect(s, x, y, Inches(5.45), Inches(2.05), fill=CARD, line=HAIR, lw=1.0, rounded=True, shadow=True)
    rect(s, x, y+Inches(0.3), Inches(0.09), Inches(1.45), fill=c)
    txt(s, x+Inches(0.45), y+Inches(0.3), Inches(4.8), Inches(0.6), h, size=21, color=INK, bold=True, font=BODYF)
    txt(s, x+Inches(0.45), y+Inches(1.0), Inches(4.8), Inches(1.0), b, size=13.5, color=BODY, spacing=1.2)
page_num(s, 3)

# ===================================================================
# SLIDE 4 - NODE LIBRARY
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, Inches(0.9), Inches(0.7), "The Node Library")
txt(s, Inches(0.85), Inches(1.25), Inches(11.6), Inches(1.0),
    "20 executable nodes. 5 categories. One registry.", size=34, color=INK, bold=True, font=HEAD)
cats = [
    ("TRIGGERS", C_TRIG, "Manual\nWebhook\nSchedule\nInterval"),
    ("ACTIONS", C_ACT, "HTTP\nEmail\nLog\nDelay\nWrite File"),
    ("LOGIC", C_LOG, "If / Else\nSwitch\nFilter\nMerge\nLoop"),
    ("DATA", C_DAT, "Set\nTransform\nCode\nTemplate\nMath"),
    ("AI", C_AI, "Generate\n\u2014 summarize\n\u2014 classify\n\u2014 extract"),
]
cw = Inches(2.3); gap = Inches(0.12); x0 = Inches(0.9); y = Inches(2.65)
for i, (name, c, body) in enumerate(cats):
    x = x0 + i * (cw + gap)
    rect(s, x, y, cw, Inches(3.1), fill=CARD, line=HAIR, lw=1.0, rounded=True, shadow=True)
    rect(s, x, y, cw, Inches(0.62), fill=c, rounded=True)
    rect(s, x, y+Inches(0.3), cw, Inches(0.32), fill=c)  # square off bottom of header
    txt(s, x, y, cw, Inches(0.62), name, size=14.5, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=BODYF, tracking=120)
    txt(s, x, y+Inches(0.85), cw, Inches(2.1), body, size=14, color=BODY,
        align=PP_ALIGN.CENTER, spacing=1.3)
txt(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(0.8),
    "Adding a node = one declarative definition + one executor. The palette, inspector, validation, "
    "command palette, and engine all pick it up automatically.", size=13.5, color=MUTE, spacing=1.2, italic=True)
page_num(s, 4)

# ===================================================================
# SLIDE 5 - ENGINE (light theme, matching)
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, Inches(0.9), Inches(0.7), "The Engine")
txt(s, Inches(0.85), Inches(1.25), Inches(11.6), Inches(1.0),
    "A real DAG runner \u2014 built from scratch.", size=36, color=INK, bold=True, font=HEAD)
# left card: pipeline
rect(s, Inches(0.9), Inches(2.65), Inches(5.75), Inches(3.85), fill=CARD, line=HAIR, lw=1.0, rounded=True, shadow=True)
steps = [
    ("1  Analyze", "Kahn's topological sort with cycle detection \u2014 invalid graphs never run."),
    ("2  Route", "Edge-activation branch routing; untaken branches are genuinely skipped."),
    ("3  Execute", "Nodes run in order: real HTTP, templating, sandboxed code, and AI."),
    ("4  Observe", "Live status, logs, and per-node I/O stream to the UI and run history."),
]
y = Inches(2.95)
for h, b in steps:
    txt(s, Inches(1.25), y, Inches(5.2), Inches(0.45), h, size=17, color=ACCENT, bold=True, font=MONO)
    txt(s, Inches(1.25), y+Inches(0.42), Inches(5.1), Inches(0.5), b, size=13, color=BODY, spacing=1.15)
    y += Inches(0.88)
# right card: differentiators
rect(s, Inches(6.95), Inches(2.65), Inches(5.4), Inches(3.85), fill=SOFTB, line=HAIR, lw=1.0, rounded=True)
rect(s, Inches(6.95), Inches(2.65), Inches(0.1), Inches(3.85), fill=ACCENT)
txt(s, Inches(7.35), Inches(2.95), Inches(4.7), Inches(0.5), "One engine, two hosts", size=19, color=INK, bold=True, font=BODYF)
diffs = [
    "Identical semantics in-browser and headless via API / webhooks.",
    "\"Run from here\" executes any node and its downstream subtree.",
    "An {{expression}} language: eq, gt, contains, truthy, and more.",
    "Error containment \u2014 one failing node never crashes the run.",
]
y = Inches(3.65)
for d in diffs:
    rect(s, Inches(7.35), y+Inches(0.06), Inches(0.12), Inches(0.12), fill=ACCENT2, rounded=True)
    txt(s, Inches(7.62), y, Inches(4.55), Inches(0.75), d, size=13, color=BODY, spacing=1.15)
    y += Inches(0.68)
page_num(s, 5)

# ===================================================================
# SLIDE 6 - PRO FEATURES
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, Inches(0.9), Inches(0.7), "Beyond The Basics", color=C_DAT)
txt(s, Inches(0.85), Inches(1.25), Inches(11.6), Inches(1.0),
    "Depth in every corner.", size=38, color=INK, bold=True, font=HEAD)
grid = [
    ("Auto-layout", "A layered DAG algorithm tidies the whole graph with a single keystroke."),
    ("Version snapshots", "Capture and restore local checkpoints of any workflow, instantly."),
    ("Copy / cut / paste", "Full keyboard-driven node editing with a 50-step undo & redo history."),
    ("Light & dark themes", "A token-based design system, applied before the first paint."),
    ("Command palette", "Press Cmd-K to navigate, insert any node, or run any action."),
    ("Cinematic 3D interface", "A WebGL hero, guided onboarding tour, and graceful degradation."),
]
for i, (h, b) in enumerate(grid):
    col, row = i % 3, i // 3
    x = Inches(0.9) + col * Inches(4.05)
    y = Inches(2.7) + row * Inches(2.0)
    rect(s, x, y, Inches(3.8), Inches(1.75), fill=CARD, line=HAIR, lw=1.0, rounded=True, shadow=True)
    rect(s, x+Inches(0.32), y+Inches(0.34), Inches(0.42), Inches(0.03), fill=ACCENT)
    txt(s, x+Inches(0.32), y+Inches(0.45), Inches(3.2), Inches(0.5), h, size=18, color=INK, bold=True, font=BODYF)
    txt(s, x+Inches(0.32), y+Inches(1.0), Inches(3.25), Inches(0.7), b, size=13, color=BODY, spacing=1.15)
page_num(s, 6)

# ===================================================================
# SLIDE 7 - WHY IT WINS
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, Inches(0.9), Inches(0.7), "Why It Wins", color=ACCENT2)
txt(s, Inches(0.85), Inches(1.25), Inches(11.6), Inches(1.0),
    "Comprehensive, honest, and real.", size=38, color=INK, bold=True, font=HEAD)
stats = [("20", "executable\nnode types"), ("5", "node\ncategories"),
         ("0", "paid API keys\nrequired"), ("100%", "offline-\ncapable")]
for i, (n, l) in enumerate(stats):
    x = Inches(0.9) + i * Inches(2.9)
    rect(s, x, Inches(2.6), Inches(2.65), Inches(1.95), fill=SOFTB, line=HAIR, lw=1.0, rounded=True)
    txt(s, x, Inches(2.75), Inches(2.65), Inches(0.95), n, size=46, color=ACCENT, bold=True,
        align=PP_ALIGN.CENTER, font=HEAD)
    txt(s, x, Inches(3.75), Inches(2.65), Inches(0.65), l, size=12.5, color=BODY,
        align=PP_ALIGN.CENTER, spacing=1.0)
rect(s, Inches(0.9), Inches(4.85), Inches(11.45), Inches(1.75), fill=NAVY, rounded=True)
rect(s, Inches(0.9), Inches(4.85), Inches(0.12), Inches(1.75), fill=ACCENT2)
txt(s, Inches(1.35), Inches(5.1), Inches(10.7), Inches(1.4),
    "Every feature is real and runs on-device: a from-scratch engine, 20 working nodes, full run "
    "observability, auto-layout, versioning, and a cinematic 3D interface \u2014 documented honestly "
    "and deployed live. No mock-ups. No vendor lock-in. A product that simply works.",
    size=15, color=RGBColor(0xDD,0xE6,0xF6), spacing=1.3, anchor=MSO_ANCHOR.MIDDLE)
page_num(s, 7)

# ===================================================================
# SLIDE 8 - CLOSE (light theme, matching)
# ===================================================================
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 0, Inches(0.18), EMU_H, fill=ACCENT)
rect(s, Inches(0.18), 0, Inches(0.06), EMU_H, fill=ACCENT2)
txt(s, Inches(1.1), Inches(2.3), Inches(11), Inches(1.3), "FlowForge",
    size=70, color=INK, bold=True, font=HEAD)
rect(s, Inches(1.17), Inches(3.6), Inches(0.9), Inches(0.035), fill=ACCENT2)
txt(s, Inches(1.12), Inches(3.9), Inches(10.8), Inches(0.8),
    "Automation, made visible.", size=27, color=BODY, italic=True, font=HEAD)
txt(s, Inches(1.12), Inches(4.85), Inches(10.8), Inches(0.5),
    "Now \u2014 let's watch it run, live.", size=19, color=ACCENT, bold=True, font=BODYF)
txt(s, Inches(1.12), Inches(6.4), Inches(11), Inches(0.4),
    "github.com/pal-123456789/flowforge      ·      flowforge-lac-eight.vercel.app",
    size=12.5, color=MUTE, font=MONO)

prs.save("FlowForge_Deck.pptx")
print("Saved FlowForge_Deck.pptx  (8 slides, professional theme)")
